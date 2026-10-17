import argparse
import json
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


CODE_COLUMNS = {
    "code": "Code",
    "definition": "Definition",
    "inclusion_criteria": "Inclusion Criteria",
    "exclusion_criteria": "Exclusion Criteria",
    "typical_examples": "Typical Examples",
    "atypical_examples": "Atypical Examples",
    "participants": "Participants",
    "relevance_to_RQ": "Relevance to RQ",
    "notes": "Notes",
}


def cell_value(value, separator="\n"):
    if value is None:
        return ""
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, list):
        return separator.join(str(cell_value(item)) for item in value)
    return value


def codebook_rows(codebook):
    if isinstance(codebook, dict):
        codebook = codebook["codebook"]
    return [
        {
            title: cell_value(item.get(key, ""), ", " if key == "participants" else "\n")
            for key, title in CODE_COLUMNS.items()
        }
        for item in codebook
    ]


def add_sheet(workbook, title, rows, columns=None):
    worksheet = workbook.create_sheet(title)
    columns = list(columns) if columns is not None else list(
        dict.fromkeys(key for row in rows for key in row)
    )
    worksheet.append(columns)
    for row in rows:
        values = [cell_value(row.get(column, "")) for column in columns]
        worksheet.append(values)
        for cell, value in zip(worksheet[worksheet.max_row], values):
            if isinstance(value, str):
                cell.data_type = "s"
    return worksheet


def save_workbook(workbook, output_file):
    border = Border(**{edge: Side(style="thin") for edge in ("left", "right", "top", "bottom")})
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="366092")
    for worksheet in workbook:
        worksheet.freeze_panes = "A2"
        for row in worksheet:
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        for index, column in enumerate(worksheet.columns, 1):
            longest = max(
                (len(line) for cell in column if cell.value is not None
                 for line in str(cell.value).splitlines()),
                default=0,
            )
            worksheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 18), 80)
        for row in worksheet.iter_rows(min_row=2):
            lines = max(
                (sum(max(1, ceil(len(line) / worksheet.column_dimensions[get_column_letter(cell.column)].width))
                     for line in str(cell.value).splitlines())
                 for cell in row if cell.value is not None),
                default=1,
            )
            worksheet.row_dimensions[row[0].row].height = min(max(lines * 18, 30), 300)

    output_file = Path(output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    return output_file


def convert_baseline(input_file, output_file=None):
    input_file = Path(input_file)
    data = json.loads(input_file.read_text(encoding="utf-8-sig"))
    codebook = data["Codebook"] if "Codebook" in data else data["codebook"]
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Codebook", codebook_rows(codebook), CODE_COLUMNS.values())
    return save_workbook(workbook, output_file or input_file.with_suffix(".xlsx"))


def main():
    parser = argparse.ArgumentParser(description="Convert a baseline codebook JSON to Excel.")
    parser.add_argument("-i", "--input", required=True, help="Input JSON path.")
    parser.add_argument("-o", "--output", help="Output XLSX path; defaults to beside the input JSON.")
    args = parser.parse_args()
    output_file = convert_baseline(args.input, args.output)
    print(f"Excel saved: {output_file}")


if __name__ == "__main__":
    main()
